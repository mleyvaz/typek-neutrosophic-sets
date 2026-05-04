"""Agrega apéndices con prompts y estímulos al Word del paper Type-k."""
from docx import Document
from docx.shared import Pt, Cm

doc = Document(
    "C:/Users/HP/Documents/TypeK_Neutrosophic_Paper/"
    "Smarandache_LeyvaVazquez_2026_TypeK_NS_v01.docx"
)

def heading(text, size=12, bold=True, sb=10, sa=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(sb)
    p.paragraph_format.space_after  = Pt(sa)
    r = p.add_run(text)
    r.bold = bold; r.font.size = Pt(size); r.font.name = "Times New Roman"

def body(text, size=10, indent=0, sa=4, mono=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(sa)
    p.paragraph_format.left_indent = Cm(indent)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.name = "Courier New" if mono else "Times New Roman"

def simple_table(headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        c.paragraphs[0].runs[0].bold = True
        c.paragraphs[0].runs[0].font.size = Pt(9)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.rows[ri + 1].cells[ci]
            c.text = str(val)
            c.paragraphs[0].runs[0].font.size = Pt(9)
    doc.add_paragraph()

# ── PAGE BREAK ────────────────────────────────────────────────────
doc.add_page_break()

# ════════════════════════════════════════════════════════════
# APPENDIX A — STIMULI
# ════════════════════════════════════════════════════════════
heading("Appendix A — Epistemic Stimuli", size=12)

body(
    "Table A1. Five epistemic phenomena used across all protocols. "
    "Stimuli were held constant across all vendors and repetitions.",
    size=10, sa=6,
)
simple_table(
    headers=["Phenomenon", "Statement"],
    rows=[
        ["Paradox (Logical)",
         "This sentence is false."],
        ["Ignorance (Epistemic)",
         "The number of stars in the universe is even."],
        ["Vagueness (Fuzzy)",
         "John is 1.75 meters tall, therefore John is tall."],
        ["Contradiction (Ethical)",
         "Lying to save an innocent life is morally right "
         "and wrong at the same time."],
        ["Contingency (Future)",
         "It will rain in New York tomorrow."],
    ],
)

body(
    "Table A2. Three tautology controls (Mason 2026 critique: "
    "a valid protocol should return T~1, I~0, F~0 for logical truths). "
    "All 6 vendors returned 0% Type-2 on all tautology controls.",
    size=10, sa=6,
)
simple_table(
    headers=["Control type", "Statement", "Expected (T, I, F)"],
    rows=[
        ["Tautology (Math)",         "2 + 2 = 4.",
         "(~1.0, ~0.0, ~0.0)"],
        ["Tautology (Definitional)", "All bachelors are unmarried.",
         "(~1.0, ~0.0, ~0.0)"],
        ["Tautology (Logical)",      "It is raining or it is not raining.",
         "(~1.0, ~0.0, ~0.0)"],
    ],
)

# ════════════════════════════════════════════════════════════
# APPENDIX B — PROMPTS
# ════════════════════════════════════════════════════════════
heading("Appendix B — Prompt Templates (verbatim)", size=12, sb=16)
body(
    "Mason (2026) raised prompt-sensitivity as a potential confound. "
    "We address it in three ways: (1) publishing all prompts verbatim below; "
    "(2) running three S1 ablation variants (Appendix C); "
    "(3) keeping stimuli identical across all vendors. "
    "The Type-2 phenomenon is absent under all S1 variants and present under "
    "all S4-O variants, across all six vendors — ruling out prompt wording "
    "as an explanation.",
    size=10, sa=8,
)

prompts = [
    (
        "B.1  S1 — Standard Neutrosophic (baseline)",
        "You are an expert in Neutrosophic Logic. You evaluate statements "
        "using three INDEPENDENT dimensions: Truth (T), Indeterminacy (I), "
        "and Falsity (F), each on [0.0, 1.0]. These dimensions are NOT "
        "constrained to sum to 1.0. A statement can be simultaneously "
        "partially true AND partially false AND partially indeterminate. "
        "Respond with ONLY a JSON object, no other text.",
        'Evaluate this statement on three independent dimensions:\n\n'
        'Statement: "{stmt}"\n\n'
        "- Truth (T): To what degree is this statement true? [0.0 to 1.0]\n"
        "- Indeterminacy (I): To what degree is the truth value unknown, "
        "undetermined, or inherently uncertain? [0.0 to 1.0]\n"
        "- Falsity (F): To what degree is this statement false? [0.0 to 1.0]\n\n"
        'T, I, and F are independent. They need NOT sum to 1.0.\n\n'
        'Respond with ONLY: {"T": <v>, "I": <v>, "F": <v>}',
    ),
    (
        "B.2  S4 — Mason (2026) declared-loss frame",
        "You are an expert in Neutrosophic Logic and epistemic honesty. "
        "You evaluate statements using three INDEPENDENT dimensions: "
        "Truth (T), Indeterminacy (I), and Falsity (F), each on [0.0, 1.0]. "
        "Crucially, you must also declare your LOSSES: what you cannot "
        "evaluate, what limits your assessment, and why your indeterminacy "
        "value is what it is. Respond with ONLY a JSON object, no other text.",
        'Evaluate this statement and declare what you cannot evaluate:\n\n'
        'Statement: "{stmt}"\n\n'
        "- Truth (T): [0.0 to 1.0]\n"
        "- Indeterminacy (I): [0.0 to 1.0]\n"
        "- Falsity (F): [0.0 to 1.0]\n"
        '- losses: list of {"what": ..., "why": ..., "severity": [0,1]}\n\n'
        "You MUST declare at least one loss.\n\n"
        'Respond with ONLY: {"T": <v>, "I": <v>, "F": <v>, '
        '"losses": [{"what": "<>", "why": "<>", "severity": <v>}, ...]}',
    ),
    (
        "B.3  S4-N — Per-attribute neutrosophic severity",
        "You are an expert in plithogenic neutrosophic logic. You evaluate "
        "statements using a per-attribute decomposition: for each attribute "
        "you identify, you provide its own (T_v, I_v, F_v) triple in [0,1]^3. "
        "Respond with ONLY a JSON object.",
        'Statement: "{stmt}"\n\n'
        "Provide:\n"
        "1. A scalar (T, I, F) as the global evaluation.\n"
        "2. A list of attributes, each with its own per-attribute triple.\n\n"
        'Respond with ONLY: {"scalar": {"T": <v>, "I": <v>, "F": <v>}, '
        '"attributes": [{"what": "<label>", "why": "<reason>", '
        '"T_v": <v>, "I_v": <v>, "F_v": <v>}, ...]}',
    ),
    (
        "B.4  S4-O.A — Overset [0, 2]",
        "You are an expert evaluator in extended Neutrosophic Logic, working "
        "in the OVERSET regime. T, I, F are three independent dimensions, each "
        "on the EXTENDED interval [0.0, 2.0]. Anchor at 1.0: "
        "baseline-calibrated truth-support. T > 1.0: stronger evidence than "
        "baseline. T < 1.0: weaker. Worker analogy (Smarandache 2016): an "
        "employee working overtime deserves membership > 1. "
        "Stay within [0.0, 2.0]. Respond with ONLY a JSON object.",
        'Evaluate on [0.0, 2.0]:\n\nStatement: "{stmt}"\n\n'
        "- T: degree of truth (1.0 = baseline)\n"
        "- I: degree of indeterminacy (1.0 = baseline)\n"
        "- F: degree of falsity (1.0 = baseline)\n\n"
        'Respond with ONLY: {"T": <v>, "I": <v>, "F": <v>}',
    ),
    (
        "B.5  S4-O.C — Offset peer-evaluation [−1, 2]",
        "You are an expert evaluator in extended Neutrosophic Logic, working "
        "in the OFFSET regime. You receive (a) a STATEMENT and (b) the OUTPUT "
        "produced by another LLM. Evaluate the trust-impact of the OTHER "
        "MODEL's output on [-1.0, 2.0]. Anchor at 1.0: baseline-correct. "
        "T,I,F > 1.0: OVERTIME (stronger than baseline). "
        "T,I,F < 0.0: ADVERSARIAL (actively misled). "
        "Worker analogy: damaging the company => membership < 0. "
        "Respond with ONLY a JSON object.",
        'Statement: "{stmt}"\n'
        'Output produced by another model: "{other}"\n\n'
        "Evaluate the OTHER MODEL's output on [-1.0, 2.0]:\n"
        "- T: trust-impact on truth\n"
        "- I: impact on indeterminacy\n"
        "- F: impact on falsity\n\n"
        "If exceeds baseline -> > 1.0. If actively misleads -> < 0.0.\n\n"
        'Respond with ONLY: {"T": <v>, "I": <v>, "F": <v>}',
    ),
]

for title, sys_p, usr_p in prompts:
    heading(title, size=10, bold=True, sb=12, sa=2)
    body("System:", size=9, sa=1)
    body(sys_p, size=9, indent=0.3, sa=4, mono=True)
    body("User:", size=9, sa=1)
    body(usr_p, size=9, indent=0.3, sa=8, mono=True)

# ════════════════════════════════════════════════════════════
# APPENDIX C — ABLATION
# ════════════════════════════════════════════════════════════
heading("Appendix C — Prompt Ablation Results", size=12, sb=16)
body(
    "Three S1 prompt variants were tested to address Mason's (2026) "
    "sensitivity concern. All variants produced 0% Type-2 responses, "
    "confirming that the absence of Type-2 under S1 is robust to "
    "prompt wording.",
    size=10, sa=6,
)
simple_table(
    headers=["Variant", "Key difference", "Phenomena", "n", "Type-2"],
    rows=[
        ["S1 (original)", "Expert framing, explicit non-sum",
         "All 5",             "300", "0 (0.0%)"],
        ["S1.v2",         "Minimal framing, implicit constraint",
         "Paradox + Ethical", "120", "0 (0.0%)"],
        ["S1.v3",         "Auditor framing, unconstrained explicit",
         "Paradox + Ethical", "120", "0 (0.0%)"],
    ],
)
body(
    "Interpretation: the dichotomy (0% under S1 variants vs. 97.3% under "
    "S4-O.A) is not a prompt artefact. It is a property of the output "
    "range allowed by the protocol, not of the framing language.",
    size=10,
)

# ── SAVE ──────────────────────────────────────────────────────────
out = (
    "C:/Users/HP/Documents/TypeK_Neutrosophic_Paper/"
    "Smarandache_LeyvaVazquez_2026_TypeK_NS_v01.docx"
)
doc.save(out)
print("Guardado:", out)
