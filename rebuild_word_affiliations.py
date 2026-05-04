"""Reconstruye el .docx con la afiliación corregida de Maikel."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import subprocess, sys

# Regenerar desde cero con build_word.py y luego add_appendix.py
subprocess.run([sys.executable, "build_word.py"], check=True)
subprocess.run([sys.executable, "add_appendix.py"], check=True)

# Abrir y corregir solo la línea de afiliación
doc = Document(
    "C:/Users/HP/Documents/TypeK_Neutrosophic_Paper/"
    "Smarandache_LeyvaVazquez_2026_TypeK_NS_v01.docx"
)

WRONG  = "Ohigins, Santiago, Chile"
RIGHT  = "Universidad Bernardo O'Higgins, Santiago, Chile"

fixed = 0
for para in doc.paragraphs:
    if WRONG in para.text:
        for run in para.runs:
            if WRONG in run.text:
                run.text = run.text.replace(WRONG, RIGHT)
                fixed += 1

# También corregir en tablas por si acaso
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                if WRONG in para.text:
                    for run in para.runs:
                        if WRONG in run.text:
                            run.text = run.text.replace(WRONG, RIGHT)
                            fixed += 1

out = (
    "C:/Users/HP/Documents/TypeK_Neutrosophic_Paper/"
    "Smarandache_LeyvaVazquez_2026_TypeK_NS_v01.docx"
)
doc.save(out)
print(f"Guardado ({fixed} reemplazos): {out}")
